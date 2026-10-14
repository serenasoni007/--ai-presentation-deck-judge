import os
from dotenv import load_dotenv
from pptx import Presentation
from google import genai

load_dotenv()

def extract_slide_text(pptx_file):
    """Extracts text content slide-by-slide from a PPTX file."""
    prs = Presentation(pptx_file)
    slide_data = []
    for i, slide in enumerate(prs.slides):
        text_content = " ".join([shape.text for shape in slide.shapes if shape.has_text_frame])
        slide_data.append({"slide_number": i + 1, "text": text_content})
    return slide_data

def evaluate_slides_with_ai(slide_data):
    """Sends slide data to Gemini for structured, sectioned feedback."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY is missing or not found in the .env file.")
        
    client = genai.Client(api_key=api_key)
    
    prompt = f"""
    Act as an expert presentation design and content judge. 
    Here is the text content extracted from a presentation deck:
    {slide_data}
    
    Evaluate the deck for flow, content structure, conciseness, and clarity.
    
    You MUST format your exact response using these exact section headings (with markdown ###):
    ### SECTION_OVERVIEW
    Give an overall score out of 100, a short executive summary, and overall impressions.
    
    ### SECTION_STRENGTHS_WEAKNESSES
    List the key strengths as bullet points, followed by weaknesses or areas for improvement as bullet points.
    
    ### SECTION_CRITIQUE
    Provide a detailed slide-by-slide constructive critique.
    """
    
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt,
    )
    return response.text